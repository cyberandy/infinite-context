#!/usr/bin/env python3
"""
Create a PPTX presentation with V4 Swiss Style GIFs as full-slide images.

Upload the resulting .pptx to Google Drive → Open with Google Slides.
GIFs will auto-animate in Google Slides presentation mode.

Usage:
    python create_slides.py
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Emu

# Slide definitions: (GIF filename, speaker notes text)
SLIDES = [
    ("slide_01.gif", "01 — Infinite Context"),
    ("slide_02.gif", "What Is Context?"),
    ("slide_03.gif", "The Quadratic Wall"),
    ("slide_04.gif", "The Context Race"),
    ("slide_05.gif", "02 — Architecture"),
    ("slide_06.gif", "Stretching the Window"),
    ("slide_07.gif", "The Ceiling"),
    ("slide_08.gif", "Two Routes"),
    ("slide_09.gif", "03 — Memory"),
    ("slide_10.gif", "Virtual Memory for LLMs"),
    ("slide_11.gif", "Give It a Search Bar"),
    ("slide_12.gif", "Enter Recursive Language Models"),
    ("slide_13.gif", "04 — Recursion"),
    ("slide_14.gif", "How RLMs Work"),
    ("slide_15.gif", "RLM-on-KG"),
    ("slide_16.gif", "Results"),
    ("slide_17.gif", "05 — The SEO Playbook"),
    ("slide_18.gif", "From Crawl-Index-Rank to Explore-Verify-Cite"),
    ("slide_19.gif", "What to Optimize"),
    ("slide_20.gif", "Structure Is the New Moat"),
]

GIFS_DIR = Path("gifs_v4")
OUTPUT = Path("infinite_context_v4.pptx")

# Standard 16:9 dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def main():
    prs = Presentation()

    # Set 16:9 slide dimensions
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Use blank layout
    blank_layout = prs.slide_layouts[6]  # Blank

    for i, (gif_name, note_text) in enumerate(SLIDES, 1):
        gif_path = GIFS_DIR / gif_name
        if not gif_path.exists():
            print(f"  ✗ Missing: {gif_path}")
            continue

        slide = prs.slides.add_slide(blank_layout)

        # Add GIF as full-slide image (position 0,0, full width/height)
        slide.shapes.add_picture(
            str(gif_path),
            left=Emu(0),
            top=Emu(0),
            width=SLIDE_WIDTH,
            height=SLIDE_HEIGHT,
        )

        # Add speaker notes
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = note_text

        size_mb = gif_path.stat().st_size / (1024 * 1024)
        print(f"  ✓ Slide {i:2d}: {gif_name} ({size_mb:.1f} MB) — {note_text}")

    prs.save(str(OUTPUT))
    size_mb = OUTPUT.stat().st_size / (1024 * 1024)
    print(f"\n✓ Saved: {OUTPUT} ({size_mb:.1f} MB)")
    print(f"\n  Next steps:")
    print(f"  1. Upload {OUTPUT} to Google Drive")
    print(f"  2. Right-click → Open with → Google Slides")
    print(f"  3. GIFs will auto-animate in present mode")


if __name__ == "__main__":
    main()
